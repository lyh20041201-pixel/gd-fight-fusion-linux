# -*- coding: utf-8 -*-
"""按学校模板生成《毕业设计开题报告表》Word 文档。

内容来源：docs/thesis/开题报告-初稿.md（本脚本内联，改内容请同步改这里）
格式依据：开题报告模板（人工智能类）-2022级.pdf

用法：
    .venv\\Scripts\\python.exe scripts\\build_proposal_docx.py
输出：
    docs/thesis/开题报告-刘宇涵.docx
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "thesis" / "开题报告-刘宇涵.docx"
FIG1 = ROOT / "docs" / "thesis" / "figures" / "图1-系统整体架构图.png"

CN = "宋体"
EN = "Times New Roman"
BODY_PT = Pt(10.5)   # 五号
SMALL_PT = Pt(9)     # 小五号
TITLE_PT = Pt(14)    # 四号
LINE_PT = Pt(20)     # 行距 20 磅


# ---------------------------------------------------------------- 基础工具

def set_run(run, size=BODY_PT, bold=False, cn=CN, en=EN):
    run.font.size = size
    run.font.bold = bold
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), cn)
    return run


def set_para(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.LEFT, line=LINE_PT, space_after=0):
    pf = p.paragraph_format
    pf.alignment = align
    if line is not None:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = line
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    if indent_chars:
        pPr = p._element.get_or_add_pPr()
        ind = pPr.get_or_add_ind()
        # firstLineChars 以 1/100 字符为单位，200 = 2 字符
        ind.set(qn("w:firstLineChars"), str(indent_chars * 100))
        ind.set(qn("w:firstLine"), str(int(indent_chars * 10.5 * 20)))
    return p


TOKEN = re.compile(r"(\*\*.*?\*\*|\[\d+\])")


def add_rich(p, text, size=BODY_PT, superscript_refs=True):
    """写入一段文本，支持 **加粗** 与 [n] 上标引用标号。"""
    for piece in TOKEN.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            set_run(p.add_run(piece[2:-2]), size=size, bold=True)
        elif superscript_refs and re.fullmatch(r"\[\d+\]", piece):
            r = set_run(p.add_run(piece), size=size)
            r.font.superscript = True
        else:
            set_run(p.add_run(piece), size=size)
    return p


# ---------------------------------------------------------------- 内容写入器

class Writer:
    """把内容按顺序追加到"开题报告内容"这个单元格里。"""

    def __init__(self, cell):
        self.cell = cell
        self._first = True

    def _p(self):
        if self._first:
            self._first = False
            return self.cell.paragraphs[0]
        return self.cell.add_paragraph()

    def h1(self, text):
        p = set_para(self._p(), indent_chars=0)
        set_run(p.add_run(text), bold=True)

    def h2(self, text):
        p = set_para(self._p(), indent_chars=2)
        set_run(p.add_run(text), bold=True)

    def body(self, text, refs=True):
        p = set_para(self._p(), indent_chars=2)
        add_rich(p, text, superscript_refs=refs)

    def flush(self, text):
        """左顶格段落（参考文献条目用）。"""
        p = set_para(self._p(), indent_chars=0)
        add_rich(p, text, superscript_refs=False)

    def formula(self, text):
        """居中公式。`_x` 记法渲染为下标，如 AP_c。"""
        p = set_para(self._p(), indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        for piece in re.split(r"(_.)", text):
            if not piece:
                continue
            if len(piece) == 2 and piece[0] == "_":
                r = set_run(p.add_run(piece[1]))
                r.font.subscript = True
            else:
                set_run(p.add_run(piece))

    def caption(self, text):
        p = set_para(self._p(), indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_run(p.add_run(text), size=SMALL_PT)

    def blank(self):
        set_para(self._p(), indent_chars=0)

    def table(self, rows, widths_cm, header=True):
        t = self.cell.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for r_i, row in enumerate(rows):
            for c_i, val in enumerate(row):
                c = t.cell(r_i, c_i)
                c.width = Cm(widths_cm[c_i])
                p = c.paragraphs[0]
                set_para(p, indent_chars=0,
                         align=WD_ALIGN_PARAGRAPH.CENTER if (header and r_i == 0) or c_i == 0
                         else WD_ALIGN_PARAGRAPH.LEFT,
                         line=None)
                set_run(p.add_run(val), size=SMALL_PT, bold=(header and r_i == 0))
        self._first = False
        return t

    def picture(self, path, width_cm):
        p = self.cell.add_paragraph()
        set_para(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER, line=None)
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        self._first = False


# ---------------------------------------------------------------- 正文内容

def fill(w: Writer) -> None:
    p0 = set_para(w._p(), indent_chars=0)
    set_run(p0.add_run("开题报告内容："), bold=True)

    # ============ 一、选题的背景与意义 ============
    w.h1("一、选题的背景与意义")
    w.h2("1.课题背景")
    w.body("教室是校园内人员密度最高、使用频次最大的场所之一，其环境质量与安全状况直接影响师生的健康与教学秩序。二氧化碳浓度长时间超标会显著降低注意力与认知表现，课后人员滞留、消防隐患等异常情况若发现不及时，则可能造成安全事故。目前高校教室管理仍以人工巡查为主，存在覆盖不全、响应滞后、记录难以追溯等问题[1]。")
    w.body("近年来物联网技术被广泛用于教室环境监测，已有研究面向智慧教室场景对温度、湿度、粉尘、二氧化碳等多传感器节点数据的融合精度进行改进[2]。然而此类系统的信息维度单一：融合始终发生在数值层面，传感器只能反映物理量的变化，无法感知教室内的人员活动与场景状态。为弥补这一缺陷，研究者将计算机视觉引入教室监控，利用目标检测与多目标跟踪技术获取人员数量与运动轨迹[3][4]。但视觉方法的输出仍停留在“数量”层面——系统能够判断画面中存在若干人员，却无法理解这些人员正在从事何种活动，因而难以区分正常教学活动与真实风险，误报率居高不下。")
    w.body("多模态大模型的发展为解决上述问题提供了新的技术路径。此类模型能够在同一次推理中联合处理图像与文本信息，具备跨模态语义理解能力[5]，可将视觉输入从“检测到若干目标”提升为“理解目标正在做什么”，并结合环境数据给出场景级研判。然而，大模型存在推理时延不确定、输出不完全可靠、依赖外部服务等固有局限，难以直接承担对实时性与确定性要求严苛的安全报警职责。")
    w.body("因此，**本课题的研究重点是：如何将多模态大模型的语义理解能力与安全监控系统的实时性、可靠性要求相结合，设计一种兼顾二者的分层协同架构**，并在此基础上实现一套可实际部署的教室环境感知与应急提醒系统。")

    w.h2("2.课题意义")
    w.body("**（1）系统用途。** 本系统面向单间教室本地部署，统一接入多路 USB 摄像头与 ESP32-C3 无线传感器节点，完成“环境感知→视觉人员检测跟踪→多摄像头区域融合→本地规则判定→声光报警→风险事件建档→多模态复核→历史回溯”的完整闭环，为教室管理人员提供实时监控、异常提醒与事后追溯能力。")
    w.body("**（2）主要功能与需求。** 系统通过环境监测模块解决“教室现在什么状况”的问题，通过视觉感知与多模态理解模块解决“教室里正在发生什么”的问题，通过规则引擎与报警模块解决“是否需要立即处置”的问题，通过事件建档与回溯模块解决“事后如何追责与改进”的问题。")
    w.body("**（3）技术先进性。** 其一，实现图像与环境数据的真正融合推理——同一帧画面在不同的传感器上下文下会得出不同结论，两种模态在推理过程中相互约束，而非各自独立分析；其二，提出分层协同架构，本地规则层保证毫秒级响应与判定确定性，多模态理解层提供语义级研判与自然语言解释，两层能力互补；其三，提出教室安全状态与系统健康状态相互独立的双维度状态模型，避免“环境指标超标”与“传感器故障”在单一状态维度下的语义混淆。")

    w.h2("3.国内外研究现状")
    w.body("**（1）环境监测与多源信息融合。** 面向智慧教室的环境监测研究主要围绕传感器组网与融合精度展开，例如通过变步长因子自适应调节与节点密度加权，提升温度、湿度、粉尘、二氧化碳等多传感器节点数据的融合优度[2]。在更广的工程领域，“多源信息融合＋风险智能预警”的技术路线已有实践，例如面向软土地铁基坑施工的多源信息融合风险智能预警研究[6]，表明多源融合对高风险场景的早期预警具有价值。但这类工作的融合对象始终是数值型数据，判定逻辑普遍停留在静态阈值比较，缺乏对持续时间、滞回、多帧确认等抗抖动机制的处理，易产生频繁误报，更无法对场景语义作出解释。")
    w.body("**（2）视觉人员检测与计数。** YOLO 系列单阶段检测器凭借速度与精度的平衡成为实时场景的主流选择[3]；在多目标跟踪方面，ByteTrack 通过对低置信度检测框的二次关联显著提升了遮挡场景下的轨迹保持能力[4]。在课堂场景中，已有研究构建了包含 5570 张图像的课堂行为数据库，并通过改进 YOLOv5 的损失函数实现举手、站立、使用手机等行为的识别，mAP@0.5 达到 96.3%[7]，表明视觉方法在教室场景具备可用精度。但该类工作依赖为特定行为类别专门标注与训练，类别之外的情形无法覆盖；且单摄像头视角难以覆盖完整教室空间，多摄像头场景下的重复计数问题尚缺乏轻量化解决方案。")
    w.body("**（3）多模态大模型在安防与异常检测中的应用。** 随着视觉语言模型能力提升[5]，已有研究尝试将其用于监控视频的异常检测与解释。VERA 通过语言化学习使视觉语言模型能够对异常事件给出可解释的自然语言说明[8]，验证了大模型在“解释为什么异常”上的独特价值；SlowFastVAD 则受人类视觉双通路启发，先由轻量检测器给出粗粒度异常置信度，仅将少量模糊片段交由推理较慢但可解释的视觉语言模型精细分析，在保持精度的同时显著降低计算开销[9]。在工业安全领域，已有研究提出基于多模态大模型的石油化工风险预警与应急处置系统，采用“多模态感知—数据融合—认知推理—决策响应”的四层架构融合视频、声学、温度、气体浓度等多源异构数据，并报告了优于单一模态方法的检测准确率[10]；值得注意的是，该研究同时明确指出，大模型在安全关键场景的工程化部署仍面临幻觉风险控制、模型可解释性与算力成本等挑战。")
    w.body("**（4）本课题的切入点。** 综合上述现状，现有工作在“感知维度”与“语义理解”两方面各有侧重，但缺少将二者在一套面向真实报警的实时系统中协同的完整方案。本课题在文献[9]所验证的快慢协同思路基础上，进一步面向安全报警场景强化可靠性约束：以本地规则引擎保证安全链路的实时性与确定性，以多模态大模型承担语义级研判与解释，并明确要求大模型不可用时系统仍能完整完成检测、报警与建档；同时通过多摄像头区域融合与双维度状态模型，解决工程落地中的重复计数与状态语义混淆问题。")

    # ============ 二、设计目标 ============
    w.h1("二、设计目标")
    w.body("系统拟实现以下四项目标：")
    w.body("**（1）多模态融合感知。** 接入温湿度、二氧化碳、光照、噪声、烟感等环境传感器与多路 USB 摄像头，实现图像、环境数据与规则依据文本的融合分析。风险事件触发时，将关键帧图像、事发时刻传感器快照、人员轨迹摘要与规则判定依据一并送入多模态大模型，输出事件性质研判、置信度、建议处置措施与是否需要人工介入的结论。")
    w.body("**（2）智能分析与应急提醒。** 实现支持阈值、持续时长、N 帧中 M 帧、滑动窗口中位数、滞回阈值与冷却时间的规则引擎，覆盖烟感触发、二氧化碳严重超标、噪声持续异常、教室超员、课后人员滞留、多摄像头人数矛盾、摄像头遮挡、关键设备离线等风险类型；触发后驱动 RGB 警示灯与蜂鸣器实现声光提醒，并完成事件建档与关键帧留存。")
    w.body("**（3）可视化 Web 管理系统。** 开发实时监控、设备状态、报警中心、历史回溯、数据趋势与系统设置六个功能页面，实现环境状态实时展示、多路视频取流、历史数据查询与预警管理，全部阈值与规则参数支持在线修改。")
    w.body("**（4）系统可靠性。** 多模态大模型不可用、超时或返回非法结果时，本地检测、判定、报警与建档链路不受任何影响；未接入的设备明确显示为离线，不以模拟数据替代真实数据。")

    # ============ 三、设计思路 ============
    w.h1("三、设计思路")
    w.h2("1.研究方法")
    w.body("系统涉及的人工智能技术及其基本原理如下。")
    w.body("**目标检测。** 采用 YOLO 系列单阶段检测器，将检测建模为回归问题，主干网络提取多尺度特征，检测头在特征图上直接预测边界框坐标与类别置信度，经非极大值抑制后输出结果。本系统只保留 person 类别，采用预训练权重，不进行自训练。")
    w.body("**多目标跟踪。** 采用 ByteTrack 算法，先用高置信度检测框与已有轨迹进行第一次匈牙利匹配，再用低置信度检测框对未匹配轨迹进行第二次关联，从而在人员相互遮挡时保持轨迹编号连续。系统只使用匿名轨迹编号，不做人脸识别，不记录身份信息。")
    w.body("**多模态跨模态推理。** 将关键帧图像编码后与结构化文本上下文组织为同一次推理的输入，由多模态大模型在统一表示空间中联合处理两种模态，输出符合预定义 JSON Schema 的结构化结论。图像提供场景语义，文本提供环境约束，二者相互制约共同决定推理结果。")

    w.h2("2.技术路线")
    w.body("**（1）数据集介绍**")
    w.body("系统使用两类数据集。")
    w.body("检测器评测集：采用 COCO val2017 验证集（5000 张图像）的 person 类别标注，用于评估预训练检测器在本任务上的基础性能。该数据集为公开数据集，标注完备，类别定义与本系统需求一致，且与预训练权重的训练集（train2017）不重叠，可避免在训练数据上自测导致的指标虚高。该项评测已完成，结果见表 1。")
    w.body("多模态复核样本集：由本系统在运行过程中自动生成，每条样本包含事发前、事发时、事发后三组关键帧图像，事发时刻的全部传感器读数、逐路与全局人数、人员轨迹摘要、触发的规则及其判定依据，构成一条完整的图文多模态记录。人工只需为每条样本标注“是否为真实风险”的二值标签作为评价基准。计划采集不少于 200 条样本，按 7:3 划分为调试集与评测集，覆盖全部已实现的风险事件类型。")

    w.body("**（2）数据预处理**")
    w.body("图像侧：从 5 秒环形缓冲中按时间偏移抽取关键帧，统一进行 letterbox 等比缩放与归一化以适配检测器输入尺寸；送入多模态大模型前压缩为 JPEG 并进行 base64 编码，按配置上限限制单次请求的帧数以控制 token 消耗。")
    w.body("文本侧：对传感器读数进行有效性校验与单位归一，剔除因丢包或校验失败产生的无效值；将传感器快照、人数统计、轨迹摘要与规则依据组织为结构化字典，按固定模板装配为提示词；通过 JSON Schema 约束模型输出格式，非法输出直接判为失败并降级处理，不进入后续流程。")

    w.body("**（3）特征提取与模型设计**")
    w.body("系统的感知与推理链路分为三级。第一级由 YOLO 检测器从视频帧中提取人员目标的空间特征，输出边界框与置信度；第二级由 ByteTrack 在时间维度上关联检测结果，提取人员的轨迹特征，包括停留时长、位移与数量变化趋势；第三级将上述视觉特征连同环境传感器特征装配为多模态提示，由大模型完成跨模态推理。")
    w.body("需要强调的是，**风险目标的坐标只能来源于检测器与跟踪器，多模态大模型不参与任何坐标生成**，其输出仅包含事件性质研判、证据描述、建议动作与人工介入建议。")

    w.body("**（4）模型评价算法**")
    w.body("检测器评价。以交并比衡量预测框与真值框的重合程度：")
    w.formula("IoU = |A ∩ B| / |A ∪ B|")
    w.body("在给定 IoU 阈值下统计真正例 TP、假正例 FP 与假负例 FN，计算查准率 P、查全率 R 与 F1 值：")
    w.formula("P = TP / (TP + FP)")
    w.formula("R = TP / (TP + FN)")
    w.formula("F1 = 2PR / (P + R)")
    w.body("对查准率-查全率曲线积分得到平均精度 AP，并对类别取均值得到 mAP：")
    w.formula("AP = ∫₀¹ P(r) dr")
    w.formula("mAP = (1/C) Σ AP_c")
    w.body("主要报告 IoU 阈值为 0.5 时的 mAP@0.5 与 0.5:0.95 区间的 mAP，并在部署环境下实测每秒处理帧数 FPS 以评估实时性。目前已完成检测器的基础性能评测，评测在 COCO val2017 全量验证集上进行，推理设备为 CPU，结果见表 1。")
    w.caption("表 1  预训练检测器 person 类别评测结果")
    w.table([
        ["指标", "数值"],
        ["查准率 Precision", "0.7566"],
        ["查全率 Recall", "0.6709"],
        ["F1 值", "0.7112"],
        ["mAP@0.5", "0.7416"],
        ["mAP@0.5:0.95", "0.5084"],
        ["单帧推理耗时（640×480）", "31.1 ms"],
        ["平均帧率", "32.20 FPS"],
    ], widths_cm=[8.0, 8.0])
    w.body("同批评测得到的全类别 mAP@0.5:0.95 为 0.3681，与 Ultralytics 官方公布的 YOLOv8n 基准值基本一致，表明评测流程配置正确。person 类别指标显著高于全类别平均水平，说明预训练权重对人员目标的检测能力可以满足本系统需求，无需自行训练。")
    w.body("由实测帧率可推算：单进程串行处理 N 路摄像头时每路约可达 32.20/N FPS，四路并发时每路约 8 FPS，低于摄像头 15 FPS 的采集速率。这一结果从实测层面印证了系统采用有界帧缓冲（处理速度不足时丢弃旧帧而非无限堆积）这一设计的必要性。")
    w.body("跟踪器评价。采用多目标跟踪准确度 MOTA 与身份一致性指标 IDF1：")
    w.formula("MOTA = 1 − (FN + FP + IDSW) / GT")
    w.formula("IDF1 = 2·IDTP / (2·IDTP + IDFP + IDFN)")
    w.body("其中 IDSW 为身份切换次数，GT 为真值目标总数。")
    w.body("多模态复核评价。以人工标注的“是否为真实风险”为基准，统计复核结论的准确率 Acc：")
    w.formula("Acc = (TP + TN) / (TP + TN + FP + FN)")
    w.body("并重点考察误报识别率，即在实际为误报的样本中被模型正确判定为误报的比例，以衡量多模态复核对系统整体误报率的抑制效果。同时记录单次复核的响应时延与 token 消耗，评估其工程可用性。")

    w.h2("3.系统设计")
    w.body("**（1）系统功能需求分析**")
    w.body("系统划分为六个功能模块，各模块需求如表 2 所示。")
    w.caption("表 2  系统功能需求分析表")
    w.table([
        ["序号", "功能模块", "主要功能", "需求分析"],
        ["1", "环境监测", "采集温湿度、CO₂、光照、噪声、烟感数据；节点在线状态与丢包统计",
         "需在节点掉线、丢包、序号跳变时准确判定设备状态，不得将无数据误显示为正常"],
        ["2", "视觉感知", "多路摄像头采集、人员检测、轨迹跟踪、标注渲染、关键帧缓冲",
         "需支持多路并发与实时处理，检测能力不可用时应停用而非输出虚假结果"],
        ["3", "区域融合", "ROI 判定、多摄像头人数融合、人数矛盾检测",
         "需解决重叠视野下的重复计数问题，并在部分摄像头离线时给出可信的降级结果"],
        ["4", "智能判定", "规则引擎评估、安全状态维护、事件去重与自动解除",
         "需支持持续时长、滞回、冷却等抗抖动机制；硬报警须立即执行"],
        ["5", "应急提醒", "RGB 警示灯与蜂鸣器控制、指令确认与重试、人工静音与确认",
         "需保证指令可达并留有执行审计记录；模拟模式下不得驱动真实执行器"],
        ["6", "多模态复核", "关键帧与环境数据装配、大模型推理、结构化结论入库与展示",
         "需异步执行且队列有界，模型不可用时不得阻塞任何本地链路"],
    ], widths_cm=[1.2, 2.2, 5.8, 6.8])
    w.body("系统面向三类用户角色：值班教师关注实时状态与报警提醒；管理人员关注历史回溯与事件处置；系统维护人员关注设备健康状态与参数配置。")

    w.body("**（2）系统整体设计架构**")
    w.body("系统采用五层架构，如图 1 所示。设备接入层负责串口协议解析与摄像头采集；感知处理层完成视觉流水线与多摄像头区域融合；判定决策层包含规则引擎、报警控制、事件建档与多模态复核；服务接口层提供 REST 接口与 WebSocket 实时推送；展示层为六个功能页面。各层之间单向依赖，真实设备与模拟设备实现同一抽象接口，使上层逻辑与运行模式解耦。图中多模态复核以虚线框标注并置于主处理栈之外，以体现其旁路、允许失败、不参与报警决策的定位。")
    w.picture(FIG1, 15.5)
    w.caption("图 1  系统整体架构图")
    w.body("系统主控制循环以 2 Hz 周期运行，依次执行多摄像头融合刷新、教室快照组装、规则引擎评估、判定结果处理与实时数据广播。多模态复核在独立的有界队列中异步执行，其结果只写入事件记录并更新前端展示，不参与报警决策，从而在引入外部服务能力的同时保证本地链路的确定性。")

    # ============ 四、将提交的成果 ============
    w.h1("四、将提交的成果")
    w.body("1. 毕业设计报告。")
    w.body("2. 模型和软件源代码。")
    w.body("3. 完整数据集（多模态复核样本集及其标注，公开数据集的获取脚本与划分清单）。")
    w.body("4. 系统部署和测试的视频。")

    # ============ 五、进度计划 ============
    w.h1("五、进度计划")
    w.body("2026 年 7 月—8 月　开展多模态模型、智能感知系统相关技术调研，完成系统总体方案设计，实现整体框架搭建。")
    w.body("2026 年 9 月　完成开题报告撰写并参加开题答辩，根据意见完善技术方案。")
    w.body("2026 年 10 月—12 月　完成数据采集处理、模型设计、系统开发，完成毕设报告初稿。")
    w.body("2027 年 1 月　完成中期检查，优化模型结构和系统功能。")
    w.body("2027 年 2 月—3 月　完成系统测试、论文完善，准备毕业审定检查。")

    # ============ 六、参考文献 ============
    w.h1("六、参考文献")
    for ref in [
        "[1] Fan Y, Cao X, Zhang J, et al. Short-term exposure to indoor carbon dioxide and cognitive task performance: A systematic review and meta-analysis[J]. Building and Environment, 2023, 237: 110331.",
        "[2] 王金锟. 基于物联网的智慧教室环境多传感器节点数据自适应融合方法[J]. 桂林航天工业学院学报, 2023(2): 226-231.",
        "[3] Terven J, Córdova-Esparza D M, Romero-González J A. A comprehensive review of YOLO architectures in computer vision: From YOLOv1 to YOLOv8 and YOLO-NAS[J]. Machine Learning and Knowledge Extraction, 2023, 5(4): 1680-1716.",
        "[4] Zhang Y, Sun P, Jiang Y, et al. ByteTrack: Multi-object tracking by associating every detection box[A]. 见: Avidan S, Brostow G, Cissé M, 等编. Computer Vision – ECCV 2022[C]. Cham: Springer, 2022: 1-21.",
        "[5] Yin S, Fu C, Zhao S, et al. A survey on multimodal large language models[J]. National Science Review, 2024, 11(12): nwae403.",
        "[6] 曾佳佳. 基于多源信息融合的软土地铁基坑施工风险智能预警研究[D]. 南昌: 东华理工大学, 2025.",
        "[7] 李甜甜. 基于计算机视觉的学生课堂行为识别研究与应用[D]. 太原: 太原师范学院, 2023.",
        "[8] Ye M, Liu W, He P. VERA: Explainable video anomaly detection via verbalized learning of vision-language models[A]. 见: Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition[C]. 2025.",
        "[9] Ding Z, Zhang H, Wu P, et al. SlowFastVAD: Video anomaly detection via integrating simple detector and RAG-enhanced vision-language model[EB/OL]. https://arxiv.org/abs/2504.10320, 2025-04-14.",
        "[10] 安天志, 董岳, 韩毅, 等. 基于多模态大模型的石油化工智能安全风险预警与应急处置系统研究[A]. 见: 2026中国石油石化企业信息技术交流大会论文集[C]. 2026: 51-61.",
    ]:
        w.flush(ref)


# ---------------------------------------------------------------- 组装文档

HEADER_ROWS = [
    ("题目名称", "多模态大模型教室环境感知与应急提醒系统开发"),
]
INFO2 = ["选题类型", "毕业设计", "专业", "计算机科学与技术", "指导教师", "吉鹏硕"]
INFO3 = ["学生姓名", "刘宇涵", "学号", "20231205419", "班级", "23计本2"]
COLS_CM = [2.4, 2.8, 1.8, 4.2, 2.1, 3.2]   # 合计 16.5cm


def cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0):
    p = cell.paragraphs[0]
    set_para(p, indent_chars=indent, align=align, line=None)
    set_run(p.add_run(text), bold=bold)


def main() -> int:
    doc = Document()

    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.0)

    # 默认样式
    normal = doc.styles["Normal"]
    normal.font.name = EN
    normal.font.size = BODY_PT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN)

    # 附件号 + 标题
    p = set_para(doc.add_paragraph(), indent_chars=0, line=None)
    set_run(p.add_run("附件 4-1"), size=TITLE_PT, bold=True)
    p = set_para(doc.add_paragraph(), indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER, line=None)
    set_run(p.add_run("毕业设计开题报告表"), size=TITLE_PT, bold=True)

    t = doc.add_table(rows=7, cols=6)
    t.style = "Table Grid"
    t.autofit = False
    for row in t.rows:
        for i, c in enumerate(row.cells):
            c.width = Cm(COLS_CM[i])

    # 第 1 行：题目名称
    cell_text(t.cell(0, 0), "题目名称", bold=True)
    m = t.cell(0, 1).merge(t.cell(0, 5))
    cell_text(m, HEADER_ROWS[0][1], align=WD_ALIGN_PARAGRAPH.CENTER)

    # 第 2、3 行
    for r, vals in ((1, INFO2), (2, INFO3)):
        for i, v in enumerate(vals):
            cell_text(t.cell(r, i), v, bold=(i % 2 == 0))

    # 第 4 行：开题报告内容（跨 6 列）
    content = t.cell(3, 0).merge(t.cell(3, 5))
    content.width = Cm(sum(COLS_CM))
    fill(Writer(content))

    # 第 5 行：指导教师意见
    op = t.cell(4, 0).merge(t.cell(4, 5))
    cell_text(op, "指导教师意见：", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    for txt, align in [("", WD_ALIGN_PARAGRAPH.LEFT),
                       ("", WD_ALIGN_PARAGRAPH.LEFT),
                       ("", WD_ALIGN_PARAGRAPH.LEFT),
                       ("指导教师签字：", WD_ALIGN_PARAGRAPH.RIGHT),
                       ("年　　月　　日", WD_ALIGN_PARAGRAPH.RIGHT)]:
        p = set_para(op.add_paragraph(), indent_chars=0, align=align, line=None)
        set_run(p.add_run(txt))

    # 第 6 行：开题审查小组意见
    rev = t.cell(5, 0).merge(t.cell(5, 5))
    cell_text(rev, "开题审查小组意见：", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    for txt, align in [("", WD_ALIGN_PARAGRAPH.LEFT),
                       ("", WD_ALIGN_PARAGRAPH.LEFT),
                       ("", WD_ALIGN_PARAGRAPH.LEFT),
                       ("组长签字：", WD_ALIGN_PARAGRAPH.RIGHT),
                       ("年　　月　　日", WD_ALIGN_PARAGRAPH.RIGHT)]:
        p = set_para(rev.add_paragraph(), indent_chars=0, align=align, line=None)
        set_run(p.add_run(txt))

    # 第 7 行：备注
    note = t.cell(6, 0).merge(t.cell(6, 5))
    cell_text(note, "备注：选题类型：毕业设计、毕业论文、社会调查报告、作品展示、"
                    "毕业汇报演出、其他。（2000－3000 字）",
              align=WD_ALIGN_PARAGRAPH.LEFT)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print("已生成:", OUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
